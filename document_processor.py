# ============================================================
# AI Knowledge Base SaaS — Document Processor v2
# Supports PDF + TXT + Word (.docx)
# Built by: Asfer Saeed — Day 33
# ============================================================

import os
import re
import tempfile
from typing import List

# ── Text Chunking ─────────────────────────────────────
def chunk_text(
    text: str,
    chunk_size: int = 400,
    overlap: int = 50
) -> List[str]:
    """Split text into overlapping chunks"""
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        return []

    words = text.split()
    if len(words) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        if chunk.strip():
            chunks.append(chunk)
        start += chunk_size - overlap

    return chunks

# ── Text Extractors ───────────────────────────────────
def extract_from_txt(file_path: str) -> str:
    """Extract text from TXT file"""
    try:
        with open(file_path, "r",
                  encoding="utf-8",
                  errors="ignore") as f:
            return f.read()
    except Exception as e:
        return f"Error reading TXT: {e}"

def extract_from_pdf(file_path: str) -> str:
    """Extract text from PDF file"""
    try:
        import pypdf
        text = ""
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            total_pages = len(reader.pages)
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n[Page {i+1}]\n{page_text}"
        print(f"  📄 PDF: {total_pages} pages extracted")
        return text
    except ImportError:
        return "Error: pip install pypdf"
    except Exception as e:
        return f"PDF error: {e}"

def extract_from_docx(file_path: str) -> str:
    """Extract text from Word .docx file"""
    try:
        from docx import Document
        doc = Document(file_path)
        text = ""

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ])
                if row_text:
                    text += row_text + "\n"

        print(f"  📝 Word: {len(doc.paragraphs)} paragraphs")
        return text
    except ImportError:
        return "Error: pip install python-docx"
    except Exception as e:
        return f"Word error: {e}"

# ── Main Processor ────────────────────────────────────
def process_uploaded_file(
    file_content: bytes,
    filename: str,
    chunk_size: int = 400
) -> dict:
    """
    Process uploaded file bytes:
    1. Save to temp file
    2. Extract text based on type
    3. Chunk into pieces
    4. Return metadata + chunks
    """
    ext = filename.lower().split(".")[-1]

    # Save to temp file
    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=f".{ext}"
    ) as tmp:
        tmp.write(file_content)
        tmp_path = tmp.name

    try:
        # Extract text
        if ext == "pdf":
            text = extract_from_pdf(tmp_path)
            file_type = "pdf"
        elif ext == "docx":
            text = extract_from_docx(tmp_path)
            file_type = "word"
        elif ext in ["txt", "md", "csv"]:
            text = extract_from_txt(tmp_path)
            file_type = "text"
        else:
            text = extract_from_txt(tmp_path)
            file_type = "text"

        # Chunk text
        chunks = chunk_text(text, chunk_size)

        return {
            "filename": filename,
            "file_type": file_type,
            "text": text,
            "chunks": chunks,
            "chunk_count": len(chunks),
            "word_count": len(text.split()),
            "char_count": len(text),
            "file_size": len(file_content)
        }

    finally:
        # Always clean up temp file
        try:
            os.unlink(tmp_path)
        except:
            pass

def process_text_directly(
    title: str,
    content: str
) -> dict:
    """Process text content directly"""
    chunks = chunk_text(content)
    return {
        "filename": f"{title}.txt",
        "file_type": "text",
        "text": content,
        "chunks": chunks,
        "chunk_count": len(chunks),
        "word_count": len(content.split()),
        "char_count": len(content),
        "file_size": len(content.encode())
    }